import httpx
import sys

def verify_all():
    print("=== Testing Endpoints & Verification ===")
    
    # 1. Health
    with httpx.Client(timeout=10.0) as client:
        r = client.get("http://127.0.0.1:8000/api/health")
        assert r.status_code == 200, f"Health check failed: {r.status_code}"
        print("[OK] Health Check passed:", r.json())
        
        # 2. Dashboard Stats
        r = client.get("http://127.0.0.1:8000/api/dashboard/stats")
        assert r.status_code == 200, f"Dashboard stats failed: {r.status_code}"
        stats = r.json()
        print(f"[✓] Dashboard Stats passed: {stats['total_students_in_db']} students, {stats['high_risk_count']} high risk, Accuracies: {stats['model_accuracies']}")
        
        # 3. Model Metrics
        r = client.get("http://127.0.0.1:8000/api/model-metrics")
        assert r.status_code == 200, f"Model metrics failed: {r.status_code}"
        res_json = r.json()
        metrics = res_json["models"]
        best_model = res_json.get("best_model")
        best_train_acc = res_json.get("best_training_accuracy")
        assert "Ensemble" not in metrics, "Ensemble must not be in metrics"
        assert best_model in ["Logistic Regression", "Random Forest", "XGBoost"]
        print(f"[✓] Model Metrics passed for models: {list(metrics.keys())}")
        print(f"    Selected Best Model by Training Accuracy: {best_model} ({best_train_acc * 100:.2f}%)")
        for m, d in metrics.items():
            print(f"    - {m}: Train Acc={d.get('train_accuracy', d['accuracy']) * 100:.2f}%, Test Acc={d['accuracy'] * 100:.2f}%, F1={d['f1_score'] * 100:.2f}%")
            
        # 4. Predict Single Student
        pred_payload = {
            "student": {
                "student_id": "DEMO-STU-001",
                "name": "Rohan Sharma",
                "attendance_percentage": 52.0,
                "study_hours_per_week": 5.0,
                "previous_grade": 48.0,
                "assignment_score": 45.0,
                "assessment_score": 42.0,
                "participation_score": 38.0,
                "sleep_hours": 5.0,
                "tutoring_sessions": 0,
                "extracurricular_activities": "No",
                "parental_education": "High School",
                "internet_access": "No"
            },
            "model_name": "XGBoost"
        }
        r = client.post("http://127.0.0.1:8000/api/predict", json=pred_payload)
        assert r.status_code == 200, f"Prediction failed: {r.status_code}, {r.text}"
        pred = r.json()
        print(f"[✓] Prediction passed: Class={pred['predicted_class']}, Risk={pred['risk_category']} ({pred['risk_score']}%), Confidence={pred['confidence']}")
        print(f"    Top Contributing Risk Factor: {pred['contributing_factors'][0]['description']}")
        
        # 4b. Model Consensus Verification
        r_cons = client.post("http://127.0.0.1:8000/api/predict-consensus", json=pred_payload)
        assert r_cons.status_code == 200, f"Model Consensus failed: {r_cons.status_code}"
        cons_data = r_cons.json().get("consensus", {})
        print(f"[✓] Model Consensus passed: Final={cons_data.get('final_prediction')}, Agreement={cons_data.get('models_agree_display')} ({cons_data.get('consensus_percentage')}%), Status={cons_data.get('consensus_status')}")

        # 5. LLM Explanation & Recommendations
        llm_payload = {
            "student_id": pred["student_id"],
            "student_name": pred["student_name"],
            "input_data": pred["input_data"],
            "predicted_class": pred["predicted_class"],
            "confidence": pred["confidence"],
            "risk_category": pred["risk_category"],
            "risk_score": pred["risk_score"],
            "contributing_factors": pred["contributing_factors"],
            "prediction_id": pred["prediction_id"],
            "consensus": cons_data
        }
        r = client.post("http://127.0.0.1:8000/api/llm/explain", json=llm_payload)
        assert r.status_code == 200, f"LLM Explanation failed: {r.status_code}, {r.text}"
        expl = r.json()
        print(f"[✓] LLM Explanation passed ({expl['model_used']}):")
        print(f"    Summary: {expl['explanation'][:100]}...")
        print(f"    Causes ({len(expl['possible_causes'])}): {expl['possible_causes'][0]}")
        print(f"    Recommendations ({len(expl['personalized_recommendations'])}): {expl['personalized_recommendations'][0]}")
        print(f"    Early Interventions ({len(expl['early_interventions'])}): {expl['early_interventions'][0]}")
        
        # 6. Batch Prediction with CSV file
        with open("data/raw/sample_upload_template.csv", "rb") as f:
            files = {"file": ("test_batch.csv", f, "text/csv")}
            data = {"model_name": "Random Forest"}
            r = client.post("http://127.0.0.1:8000/api/batch-predict", files=files, data=data)
            assert r.status_code == 200, f"Batch predict failed: {r.status_code}, {r.text}"
            batch = r.json()
            print(f"[✓] Batch Prediction passed: Processed {batch['total_students']} students, Risk breakdown: {batch['risk_distribution']}")
            
        # 7. Frontend Served via FastAPI
        r = client.get("http://127.0.0.1:8000/")
        assert r.status_code == 200 and "<html" in r.text.lower(), "FastAPI failed to serve frontend index.html"
        print("[✓] Production Frontend bundle served successfully by FastAPI on http://127.0.0.1:8000/")

        # 8. Frontend Served via Vite Dev Server
        try:
            r = client.get("http://localhost:5173/")
            if r.status_code == 200:
                print("[✓] Vite Dev Server responding successfully on http://localhost:5173/")
        except Exception:
            print("[INFO] Vite Dev Server not running; FastAPI serving production bundle.")

        # 9. AI Resume Analyzer & Placement Readiness Verification
        import docx
        import io
        test_doc = docx.Document()
        test_doc.add_paragraph("Akashraj Sundaram\nEmail: akashraj@example.com\nB.Tech in Computer Science and Engineering\nCGPA: 8.85 / 10")
        test_doc.add_paragraph("Technical Skills: Python, Java, SQL, React, Node.js, Spring Boot, Docker, AWS, Git, Data Structures, Algorithms")
        test_doc.add_paragraph("Projects: Cloud Microservices Platform with Docker, Spring Boot, and PostgreSQL. Reduced latency by 45%.")
        test_doc.add_paragraph("Certifications: AWS Certified Cloud Practitioner")
        test_buf = io.BytesIO()
        test_doc.save(test_buf)
        test_files = {"file": ("e2e_resume.docx", test_buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
        r_resume = client.post("http://127.0.0.1:8000/api/resume/analyze", files=test_files)
        assert r_resume.status_code == 200, f"Resume analyzer failed: {r_resume.status_code}, {r_resume.text}"
        res_data = r_resume.json()
        print(f"[✓] AI Resume Analyzer passed: Candidate={res_data['candidate_name']}, Matched Skills={res_data['skill_analysis']['total_matched']}, Readiness Score={res_data['placement_readiness']['readiness_score']}% ({res_data['placement_readiness']['readiness_level']})")
        print(f"    Top Company Match: {res_data['company_matching'][0]['company']} ({res_data['company_matching'][0]['alignment']})")
        print(f"    Personalized Plan Priority 1: {res_data['personalized_improvement_plan'][0]['title']}")

    print("\n🎉 ALL END-TO-END VERIFICATIONS PASSED WITH 100% SUCCESS!")

if __name__ == "__main__":
    verify_all()
