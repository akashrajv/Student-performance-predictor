import io
import re
from pathlib import Path
from typing import Dict, Any, Union

def clean_extracted_text(text: str) -> str:
    """Clean and normalize extracted resume text."""
    if not text:
        return ""
    # Normalize bullet points and dashes
    text = re.sub(r'[\u2022\u2023\u25E6\u2043\u2219\u00B7\u25CF\u25CB\u25AA\u25AB]', '\n- ', text)
    text = re.sub(r'[\u2013\u2014\u2015]', '-', text)
    # Remove control characters except newlines/tabs
    text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]', '', text)
    # Normalize excessive spaces
    text = re.sub(r'[ \t]+', ' ', text)
    # Normalize excessive newlines
    text = re.sub(r'\n{3,}', '\n\n', text)
    return text.strip()

def extract_text_from_pdf(file_source: Union[bytes, io.BytesIO, str, Path]) -> Dict[str, Any]:
    """
    Extract text from PDF file using pypdf.
    Handles bytes, BytesIO, or file paths.
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        return {
            "status": "error",
            "message": "pypdf library is not installed. Please install pypdf to process PDF files."
        }

    try:
        if isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
        elif isinstance(file_source, io.BytesIO):
            stream = file_source
        elif isinstance(file_source, (str, Path)):
            with open(file_source, "rb") as f:
                stream = io.BytesIO(f.read())
        else:
            return {
                "status": "error",
                "message": "Invalid file source provided for PDF extraction."
            }

        # Check for empty content
        stream.seek(0, io.SEEK_END)
        file_size = stream.tell()
        stream.seek(0)
        if file_size == 0:
            return {
                "status": "error",
                "message": "The uploaded PDF file is empty (0 bytes). Please upload a valid document."
            }

        reader = PdfReader(stream)
        num_pages = len(reader.pages)
        if num_pages == 0:
            return {
                "status": "error",
                "message": "The uploaded PDF contains no pages. Please upload a valid resume."
            }

        extracted_pages = []
        for i, page in enumerate(reader.pages):
            try:
                page_text = page.extract_text() or ""
                extracted_pages.append(page_text)
            except Exception:
                extracted_pages.append("")

        full_text = "\n".join(extracted_pages)
        cleaned = clean_extracted_text(full_text)

        if len(cleaned.strip()) < 50:
            return {
                "status": "error",
                "message": (
                    f"The resume contains insufficient readable text ({len(cleaned.strip())} characters). "
                    "Please ensure the PDF contains selectable text rather than a scanned image or empty page."
                ),
                "text": cleaned,
                "page_count": num_pages
            }

        return {
            "status": "success",
            "text": cleaned,
            "page_count": num_pages,
            "char_count": len(cleaned)
        }

    except Exception as e:
        return {
            "status": "error",
            "message": "Could not read the PDF document. The file may be password-protected, corrupted, or invalid."
        }

def extract_text_from_docx(file_source: Union[bytes, io.BytesIO, str, Path]) -> Dict[str, Any]:
    """
    Extract text from DOCX file using python-docx.
    Handles bytes, BytesIO, or file paths.
    """
    try:
        import docx
    except ImportError:
        return {
            "status": "error",
            "message": "python-docx library is not installed. Please install python-docx to process Word files."
        }

    try:
        if isinstance(file_source, bytes):
            stream = io.BytesIO(file_source)
        elif isinstance(file_source, io.BytesIO):
            stream = file_source
        elif isinstance(file_source, (str, Path)):
            with open(file_source, "rb") as f:
                stream = io.BytesIO(f.read())
        else:
            return {
                "status": "error",
                "message": "Invalid file source provided for DOCX extraction."
            }

        # Check for empty content
        stream.seek(0, io.SEEK_END)
        file_size = stream.tell()
        stream.seek(0)
        if file_size == 0:
            return {
                "status": "error",
                "message": "The uploaded DOCX file is empty (0 bytes). Please upload a valid document."
            }

        doc = docx.Document(stream)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        
        # Also extract table text if present
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    table_texts.append(row_text)

        all_text = "\n".join(paragraphs + table_texts)
        cleaned = clean_extracted_text(all_text)

        if len(cleaned.strip()) < 50:
            return {
                "status": "error",
                "message": (
                    f"The resume contains insufficient readable text ({len(cleaned.strip())} characters). "
                    "Please ensure the Word document contains readable content."
                ),
                "text": cleaned
            }

        return {
            "status": "success",
            "text": cleaned,
            "paragraph_count": len(paragraphs),
            "char_count": len(cleaned)
        }

    except Exception as e:
        return {
            "status": "error",
            "message": "Could not read the DOCX document. The file may be corrupted, password-protected, or not a valid Word document."
        }

def extract_resume_text(
    file_source: Union[bytes, io.BytesIO, str, Path],
    filename: str = ""
) -> Dict[str, Any]:
    """
    Unified entry point for resume text extraction.
    Automatically detects format from filename or content.
    Returns structured result with status, text, or user-friendly error message.
    """
    if file_source is None:
        return {
            "status": "error",
            "message": "No file uploaded. Please select and upload a PDF or DOCX resume."
        }

    if isinstance(file_source, (bytes, str)) and len(file_source) == 0:
        return {
            "status": "error",
            "message": "The uploaded file is empty (0 bytes). Please upload a valid document."
        }

    # If raw text string is passed directly (e.g. preset or text content)
    if isinstance(file_source, str) and not (Path(file_source).exists() and Path(file_source).is_file()):
        cleaned = clean_extracted_text(file_source)
        if len(cleaned.strip()) < 50:
            return {
                "status": "error",
                "message": f"Resume contains insufficient readable text ({len(cleaned.strip())} characters). Please provide a complete resume document.",
                "text": cleaned
            }
        return {
            "status": "success",
            "text": cleaned,
            "char_count": len(cleaned),
            "filename": filename or "resume_text",
            "file_type": "TEXT"
        }

    # Determine extension
    ext = ""
    if filename:
        ext = filename.lower().split(".")[-1]
    elif isinstance(file_source, (str, Path)):
        ext = str(file_source).lower().split(".")[-1]

    if ext == "pdf":
        res = extract_text_from_pdf(file_source)
    elif ext in ["docx", "doc"]:
        if ext == "doc":
            return {
                "status": "error",
                "message": "Legacy .doc format is not supported. Please save the document as .docx or .pdf."
            }
        res = extract_text_from_docx(file_source)
    else:
        # Try probing file magic or attempting PDF first, then DOCX
        if isinstance(file_source, bytes):
            is_pdf = file_source.startswith(b"%PDF")
            is_docx = file_source.startswith(b"PK\x03\x04")
        else:
            is_pdf = False
            is_docx = False

        if is_pdf:
            res = extract_text_from_pdf(file_source)
        elif is_docx:
            res = extract_text_from_docx(file_source)
        else:
            return {
                "status": "error",
                "message": f"Unsupported file format '{ext or 'unknown'}'. Please upload a PDF (.pdf) or Word (.docx) document."
            }

    if res.get("status") == "success":
        res["filename"] = filename
        res["file_type"] = ext.upper() if ext else "DOCUMENT"

    return res
